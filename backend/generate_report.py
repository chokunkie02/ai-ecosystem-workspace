import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # Page setup (Standard margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'TH Sarabun PSK' if 'TH Sarabun PSK' in doc.styles else 'Calibri'
    font.size = Pt(14)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("รายงานการส่งงาน Assignment #03\nBackend Setup, Redis ARQ, PostgreSQL & Label Studio Integration")
    run_title.bold = True
    run_title.font.size = Pt(18)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("รายวิชา AI Ecosystem Workspace")
    run_sub.font.size = Pt(14)

    # Student Info
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_after = Pt(14)
    p_info.add_run("ชื่อ-นามสกุล: ").bold = True
    p_info.add_run("chokun\n")
    p_info.add_run("รหัสนักศึกษา: ").bold = True
    p_info.add_run("6710110589\n")
    p_info.add_run("GitHub Repository: ").bold = True
    p_info.add_run("https://github.com/chokunkie02/ai-ecosystem-workspace.git")

    def add_section_header(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title_text)
        run.bold = True
        run.font.size = Pt(16)

    def add_code_block(code_text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        r_code = p.add_run(code_text)
        r_code.font.name = 'Consolas'
        r_code.font.size = Pt(10)

    def add_image_placeholder(caption_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"[ วางภาพแคปหน้าจอ: {caption_text} ]")
        r.bold = True
        r.font.size = Pt(12)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run(f"รูปที่: {caption_text}")
        r_cap.font.italic = True
        r_cap.font.size = Pt(11)

    # -------------------------------------------------------------
    # Work #1
    # -------------------------------------------------------------
    add_section_header("Work #1: Project Virtual Environment & Dependency Management")
    doc.add_paragraph("สร้าง Virtual Environment สำหรับโปรเจกต์ backend ด้วยเครื่องมือ uv และกำหนดไฟล์ .gitignore เพื่อป้องกันการอัปโหลดไฟล์ที่ไม่จำเป็นขึ้นไปยัง Git Repository:")
    
    add_code_block("""# คำสั่งเริ่มต้นโปรเจกต์ backend และสร้างไฟล์ .gitignore
cd c:\\eco\\friday
mkdir backend
cd backend
uv init

# เนื้อหาในไฟล์ .gitignore
.venv/
__pycache__/
*.pyc
.env

# คำสั่งติดตั้ง Dependencies ทั้งหมดที่ใช้ใน Assignment 3
uv add pydantic-settings arq sqlalchemy psycopg2-binary asyncpg label-studio-sdk python-dotenv python-docx""")

    # -------------------------------------------------------------
    # Work #2
    # -------------------------------------------------------------
    add_section_header("Work #2: Project Settings Management (pydantic-settings)")
    doc.add_paragraph("การตั้งค่าโปรเจกต์ผ่าน Pydantic BaseSettings เพื่ออ่านค่า Environment Variables จากไฟล์ .env:")
    
    add_code_block("""# ไฟล์ .env
REDIS_HOST=localhost
REDIS_PORT=6379
POSTGRES_HOST=localhost
POSTGRES_PORT=5433
POSTGRES_USER=postgres
POSTGRES_PASSWORD=mysecretpassword
POSTGRES_DB=student_db
LABEL_STUDIO_URL=http://localhost:8080
LABEL_STUDIO_API_KEY=eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...

# ไฟล์ core/config.py
from pydantic_settings import BaseSettings

class Settings(BaseSettings):
    redis_host: str
    redis_port: int
    postgres_host: str
    postgres_port: int
    postgres_user: str
    postgres_password: str
    postgres_db: str
    label_studio_url: str
    label_studio_api_key: str

    class Config:
        env_file = ".env"

settings = Settings()

# ไฟล์ sandbox/test_settings.py
from core.config import settings
print(settings.model_dump())""")

    doc.add_paragraph("ผลลัพธ์การรันคำสั่ง uv run python -m sandbox.test_settings:")
    add_code_block("{'redis_host': 'localhost', 'redis_port': 6379, 'postgres_host': 'localhost', 'postgres_port': 5433, 'postgres_user': 'postgres', 'postgres_password': 'mysecretpassword', 'postgres_db': 'student_db', 'label_studio_url': 'http://localhost:8080', 'label_studio_api_key': 'eyJhbGci...'}")

    add_image_placeholder("ผลลัพธ์การทดสอบอ่านค่า Settings (Work #2)")

    # -------------------------------------------------------------
    # Work #3
    # -------------------------------------------------------------
    add_section_header("Work #3: Redis Queue & ARQ Background Task Processing")
    doc.add_paragraph("การสร้างระบบประมวลผลงานฉากหลัง (Background Task Worker) เชื่อมต่อ Redis Server ผ่าน ARQ:")

    add_code_block("""# ไฟล์ worker_settings.py
from arq.connections import RedisSettings
from core.config import settings

async def simple_work(ctx, data):
    print(f"Received job data: {data}")
    return {"status": "done", "data": data}

class WorkerSettings:
    functions = [simple_work]
    redis_settings = RedisSettings(host=settings.redis_host, port=settings.redis_port)

# ไฟล์ enqueue.py
import asyncio
from arq import create_pool
from arq.connections import RedisSettings
from core.config import settings

async def main():
    redis = await create_pool(RedisSettings(host=settings.redis_host, port=settings.redis_port))
    await redis.enqueue_job("simple_work", {"msg": "hello from chokun (6710110589)"})
    print("Job enqueued!")

asyncio.run(main())""")

    doc.add_paragraph("ผลลัพธ์การรับงานที่ Terminal 1 (ARQ Worker):")
    add_code_block("""21:59:38: Starting worker for 1 functions: simple_work
21:59:38: redis_version=8.8.0 mem_usage=1.74M clients_connected=1 db_keys=1
21:59:44:   0.43s -> 7718993644e246fbbc825546e46442a4:simple_work({'msg': 'hello from chokun (6710110589)'})
Received job data: {'msg': 'hello from chokun (6710110589)'}
21:59:44:   0.00s <- 7718993644e246fbbc825546e46442a4:simple_work * {'status': 'done', 'data': {'msg': 'hello from chokun (6710110589)'}}""")

    add_image_placeholder("ผลลัพธ์ ARQ Worker รับ Job จาก Redis (Work #3)")

    # -------------------------------------------------------------
    # Work #4
    # -------------------------------------------------------------
    add_section_header("Work #4: PostgreSQL CRUD Operations with SQLAlchemy")
    doc.add_paragraph("การเชื่อมต่อฐานข้อมูล PostgreSQL ผ่าน SQLAlchemy เพื่อทำรายการ CRUD (Create Table, Insert, Update, Delete, Drop Table):")

    add_code_block("""# ไฟล์ postgres_test.py
from sqlalchemy import create_engine, text
from core.config import settings

url = f"postgresql+psycopg2://{settings.postgres_user}:{settings.postgres_password}@{settings.postgres_host}:{settings.postgres_port}/{settings.postgres_db}"
engine = create_engine(url)

def create_table():
    with engine.begin() as conn:
        conn.execute(text(\"\"\"
            CREATE TABLE IF NOT EXISTS students (
                id SERIAL PRIMARY KEY,
                student_id VARCHAR(50),
                name VARCHAR(100),
                age INT,
                major VARCHAR(100)
            )
        \"\"\"))
    print("Table created.")

def show_data(label):
    with engine.begin() as conn:
        rows = conn.execute(text("SELECT * FROM students")).fetchall()
        print(f"--- {label} ---")
        for row in rows:
            print(row)

def insert_data():
    with engine.begin() as conn:
        conn.execute(text("INSERT INTO students (student_id, name, age, major) VALUES (:sid, :n, :a, :m)"),
                     {"sid": "6710110589", "n": "chokun", "a": 21, "m": "AI Engineering"})
    print("Data inserted.")

def update_data():
    with engine.begin() as conn:
        conn.execute(text("UPDATE students SET age = :a WHERE name = :n"), {"a": 22, "n": "chokun"})
    print("Data updated.")

def delete_data():
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM students WHERE name = :n"), {"n": "chokun"})
    print("Data deleted.")

def drop_table():
    with engine.begin() as conn:
        conn.execute(text("DROP TABLE students"))
    print("Table dropped.")

if __name__ == "__main__":
    create_table()
    insert_data()
    show_data("After Insert")
    update_data()
    show_data("After Update")
    delete_data()
    show_data("After Delete")
    drop_table()""")

    doc.add_paragraph("ผลลัพธ์การรันคำสั่ง uv run python postgres_test.py:")
    add_code_block("""Table created.
Data inserted.
--- After Insert ---
(1, '6710110589', 'chokun', 21, 'AI Engineering')
Data updated.
--- After Update ---
(1, '6710110589', 'chokun', 22, 'AI Engineering')
Data deleted.
--- After Delete ---
Table dropped.""")

    add_image_placeholder("ผลลัพธ์การทำงาน PostgreSQL CRUD (Work #4)")

    # -------------------------------------------------------------
    # Work #5
    # -------------------------------------------------------------
    add_section_header("Work #5: Label Studio Project Setup & Python SDK Integration")
    doc.add_paragraph("การตั้งค่าโปรเจกต์ Assignment3-Example บน Label Studio อัปโหลดชุดข้อมูล Text Classification 5 ประโยค และเชื่อมต่อดึงข้อมูลผ่าน Python SDK:")

    add_code_block("""# ไฟล์ label_studio_test.py
from label_studio_sdk.client import LabelStudio
from core.config import settings

ls = LabelStudio(base_url=settings.label_studio_url, api_key=settings.label_studio_api_key)

print("--- Projects ---")
projects = list(ls.projects.list())
for p in projects:
    print(p.id, p.title)

print("\\n--- Tasks in first project ---")
if projects:
    project_id = projects[0].id
    tasks = list(ls.tasks.list(project=project_id))
    for t in tasks:
        print(t.id, t.data)
else:
    print("No projects found.")""")

    doc.add_paragraph("ผลลัพธ์การรันคำสั่ง uv run python label_studio_test.py:")
    add_code_block("""--- Projects ---
1 Assignment3-Example

--- Tasks in first project ---
1 {'text': 'This product is amazing and I love using it!'}
2 {'text': 'The quality is okay, but the delivery was slow.'}
3 {'text': 'Terrible service, I am very disappointed.'}
4 {'text': 'The application works as expected with good features.'}
5 {'text': 'Not bad, but could be improved in future updates.'}""")

    add_image_placeholder("หน้าต่าง Web UI ของ Label Studio โปรเจกต์ Assignment3-Example (Work #5)")
    add_image_placeholder("ผลลัพธ์การรัน label_studio_test.py ผ่าน Python SDK (Work #5)")

    # Save destinations
    downloads_dir = os.path.expanduser("~/Downloads")
    filepath_downloads = os.path.join(downloads_dir, "Assignment3_Report_6710110589.docx")
    filepath_local = os.path.join(os.getcwd(), "Assignment3_Report_6710110589.docx")

    try:
        doc.save(filepath_downloads)
        print(f"Report successfully saved to: {filepath_downloads}")
    except PermissionError:
        filepath_alt = os.path.join(downloads_dir, "Assignment3_Report_6710110589_v2.docx")
        doc.save(filepath_alt)
        print(f"Original file was open. Saved to alternate path: {filepath_alt}")

    doc.save(filepath_local)
    print(f"Backup copy saved to: {filepath_local}")

if __name__ == "__main__":
    create_report()
