"""
Word Document Builder Helper Module.
Provides styling and structure helper functions using python-docx with TH Sarabun PSK font support.
"""

from typing import List, Optional
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_NAME = "TH Sarabun PSK"
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)


def set_cell_background(cell, hex_color: str):
    """Set background shading color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Set internal padding (margins) of a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    """Apply soft border style to table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)


def init_document() -> docx.Document:
    """Initialize Document with standard A4 setup, 1-inch margins, and TH Sarabun PSK style defaults."""
    doc = docx.Document()

    # Page Setup - A4
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set Default Normal Style
    normal_style = doc.styles['Normal']
    font = normal_style.font
    font.name = FONT_NAME
    font.size = Pt(16)
    font.color.rgb = COLOR_BLACK

    # Set East Asia font for Thai compatibility
    rPr = normal_style._element.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:eastAsia="{FONT_NAME}" w:cs="{FONT_NAME}"/>')
    rPr.append(rFonts)

    return doc


def add_doc_title(doc: docx.Document, title_text: str, subtitle_text: str = ""):
    """Add Main Document Title banner at top of first page (No Cover Page requirement)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)

    run = p.add_run(title_text)
    run.font.name = FONT_NAME
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK

    if subtitle_text:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(18)
        run_sub = p_sub.add_run(subtitle_text)
        run_sub.font.name = FONT_NAME
        run_sub.font.size = Pt(16)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COLOR_BLACK

    # Add subtle horizontal line separator (black)
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="000000"/></w:pBdr>')
    p_div._element.get_or_add_pPr().append(p_div_border)


def add_heading_1(doc: docx.Document, text: str):
    """Add Level 1 Chapter Heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True

    # Bottom border for H1 (black)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="000000"/></w:pBdr>')
    p._element.get_or_add_pPr().append(pBdr)

    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK


def add_heading_2(doc: docx.Document, text: str):
    """Add Level 2 Sub-section Heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK


def add_heading_3(doc: docx.Document, text: str):
    """Add Level 3 Sub-section Heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK


def add_body_p(doc: docx.Document, text: str, bold_prefix: Optional[str] = None, space_after: float = 4):
    """Add body paragraph with optional bold lead-in prefix."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(16)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK

    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_BLACK
    return p


def add_bullet_p(doc: docx.Document, text: str, bold_prefix: Optional[str] = None):
    """Add bullet item paragraph."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(16)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK

    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_BLACK
    return p


def add_code_block(doc: docx.Document, code_text: str):
    """Add code block styled inside a shaded container table."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Apply thin full gray border (#CCCCCC)
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0

    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_BLACK

    # Empty paragraph after table for spacing
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)


def add_callout(doc: docx.Document, text: str, title: str = "คำอธิบายเพิ่มเติม"):
    """Add callout note box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    r_title = p.add_run(f"{title}: ")
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLACK

    r_text = p.add_run(text)
    r_text.font.name = FONT_NAME
    r_text.font.size = Pt(16)
    r_text.font.color.rgb = COLOR_BLACK

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)


def create_styled_table(
    doc: docx.Document,
    headers: List[str],
    rows: List[List[str]],
    col_widths: Optional[List[float]] = None,
):
    """Create a clean table with neutral light gray headers (#E6E6E6), white rows (#FFFFFF), and gray borders (#B0B0B0)."""
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    set_table_borders(table, color="B0B0B0", sz="4", val="single")

    # Format Header Row (#E6E6E6 background, bold black text)
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "E6E6E6")
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = FONT_NAME
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = COLOR_BLACK

    # Format Data Rows (Clean white background #FFFFFF, black text)
    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_background(row_cells[col_idx], "FFFFFF")
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15

            # Align center for short code/method/status, left for text
            if col_idx in [0, 2] and len(cell_value) < 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(15)
                r.font.color.rgb = COLOR_BLACK

    # Set Column Widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)


def add_image_placeholder(doc: docx.Document, title: str, instruction: str = ""):
    """Add a dashed visual placeholder frame for manual screenshot insertion."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="CCCCCC"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="CCCCCC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    r_icon = p.add_run("[กรอบสำหรับวางภาพแคปหน้าจอ]: ")
    r_icon.font.name = FONT_NAME
    r_icon.font.size = Pt(16)
    r_icon.font.bold = True
    r_icon.font.color.rgb = COLOR_BLACK

    r_title = p.add_run(title)
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLACK

    if instruction:
        p_ins = cell.add_paragraph()
        p_ins.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ins.paragraph_format.space_before = Pt(2)
        p_ins.paragraph_format.space_after = Pt(4)
        r_ins = p_ins.add_run(instruction)
        r_ins.font.name = FONT_NAME
        r_ins.font.size = Pt(14)
        r_ins.font.italic = True
        r_ins.font.color.rgb = COLOR_BLACK

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)


def add_image(doc: docx.Document, image_path: str, width_inches: float = 6.0):
    """Add an image centered in paragraph with specified width in inches."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    return p


def add_caption_p(doc: docx.Document, caption_text: str):
    """Add a centered figure caption paragraph underneath images/placeholders."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)

    run = p.add_run(caption_text)
    run.font.name = FONT_NAME
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p





