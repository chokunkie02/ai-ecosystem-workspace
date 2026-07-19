import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx():
    doc = docx.Document()
    
    # Page setup
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Styles Setup
    styles = doc.styles
    
    # Document Title Style
    title_style = styles.add_style('ReportTitle', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    title_font = title_style.font
    title_font.name = 'Cordia New'
    title_font.size = Pt(28)
    title_font.bold = True
    title_font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    
    # Heading 1 Style
    h1_style = styles.add_style('ReportH1', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    h1_font = h1_style.font
    h1_font.name = 'Cordia New'
    h1_font.size = Pt(20)
    h1_font.bold = True
    h1_font.color.rgb = RGBColor(79, 70, 229) # Indigo 600
    
    # Heading 2 Style
    h2_style = styles.add_style('ReportH2', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    h2_font = h2_style.font
    h2_font.name = 'Cordia New'
    h2_font.size = Pt(16)
    h2_font.bold = True
    h2_font.color.rgb = RGBColor(100, 116, 139) # Slate 500
    
    # Body Text Style
    body_style = styles['Normal']
    body_font = body_style.font
    body_font.name = 'Cordia New'
    body_font.size = Pt(14)
    body_font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    
    # 1. Document Header
    p = doc.add_paragraph('รายงานการติดตั้งและทดสอบระบบ\nPostgreSQL & Label Studio (Docker Compose)', style='ReportTitle')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run('จัดทำโดย: โชกุน (Chokun)\nลิงก์โครงการ GitHub: https://github.com/chokunkie02/ai-ecosystem-workspace')
    run_meta.font.size = Pt(12)
    run_meta.italic = True
    
    doc.add_paragraph('\n' + '='*50 + '\n').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 2. Section 1: PostgreSQL
    doc.add_paragraph('1. ระบบฐานข้อมูล PostgreSQL', style='ReportH1')
    
    doc.add_paragraph('1.1 การกำหนดค่า docker-compose (compose.yml)')
    doc.add_paragraph('เราได้กำหนดบริการ PostgreSQL ใน compose.yml โดยใช้ Named Volume ชื่อ "postgresql-data" เพื่อให้ข้อมูลคงอยู่ถาวร และเปิดพอร์ต 5433 บนเครื่อง Host เชื่อมต่อไปยังพอร์ต 5432 ในคอนเทนเนอร์ เพื่อเลี่ยงปัญหาพอร์ตชนกัน ดังนี้:')
    
    # Code block for compose.yml
    table_compose = doc.add_table(rows=1, cols=1)
    table_compose.style = 'Light Shading Accent 1'
    cell = table_compose.rows[0].cells[0]
    set_cell_background(cell, "F1F5F9")
    compose_text = """services:
  redis:
    image: redis:8.8.0-alpine
    container_name: redis
    restart: unless-stopped
    ports:
      - "6379:6379"
    volumes:
      - redis-data:/data
    command: >
      redis-server
      --appendonly yes
    healthcheck:
      test: ["CMD", "redis-cli", "ping"]
      interval: 10s
      timeout: 5s
      retries: 5

  db:
    image: postgres:15
    container_name: postgres_db
    restart: always
    environment:
      POSTGRES_USER: postgres
      POSTGRES_PASSWORD: mysecretpassword
      POSTGRES_DB: student_db
    ports:
      - "5433:5432"
    volumes:
      - postgresql-data:/var/lib/postgresql/data

  label-studio:
    image: heartexlabs/label-studio:latest
    container_name: label_studio
    restart: always
    ports:
      - "8080:8080"
    environment:
      - DJANGO_DB=default
      - POSTGRE_NAME=student_db
      - POSTGRE_USER=postgres
      - POSTGRE_PASSWORD=mysecretpassword
      - POSTGRE_PORT=5432
      - POSTGRE_HOST=db
    depends_on:
      - db

volumes:
  redis-data:
  postgresql-data:
    name: postgresql-data"""
    cell.paragraphs[0].text = compose_text
    cell.paragraphs[0].style.font.name = 'Consolas'
    cell.paragraphs[0].style.font.size = Pt(10)
    
    doc.add_paragraph('\n1.2 ชุดคำสั่ง SQL และผลลัพธ์การทดสอบตาราง students')
    doc.add_paragraph('ได้ทำการรันสคริปต์ db_test.py เพื่อเชื่อมต่อไปยังฐานข้อมูลและทดสอบคำสั่ง SQL CRUD ต่างๆ ผลลัพธ์และรันคำสั่งมีรายละเอียดดังต่อไปนี้:')
    
    # SQL Table with queries and outputs
    sql_steps = [
        ("i. คำสั่งสร้าง table", 
         "DROP TABLE IF EXISTS students;\nCREATE TABLE students (\n    id SERIAL PRIMARY KEY,\n    name VARCHAR(100) NOT NULL,\n    age INT,\n    major VARCHAR(100)\n);",
         "Table 'students' created successfully."),
        ("ii. คำสั่งเพิ่มข้อมูลเข้า table",
         "INSERT INTO students (name, age, major) VALUES\n('chokun', 21, 'วิศวะเอไอ'),\n('chompoo', 20, 'พยาบาล'),\n('jack', 22, 'วิศวะเคมี');",
         "Inserted 3 records successfully."),
        ("iii. คำสั่งแสดงข้อมูลใน table",
         "SELECT * FROM students;",
         "Results:\n(1, 'chokun', 21, 'วิศวะเอไอ')\n(2, 'chompoo', 20, 'พยาบาล')\n(3, 'jack', 22, 'วิศวะเคมี')"),
        ("iv. คำสั่งแก้ไขข้อมูลใน table",
         "UPDATE students SET age = 23 WHERE name = 'chokun';",
         "Updated 1 record(s) successfully.\nUpdated Row:\n(1, 'chokun', 23, 'วิศวะเอไอ')"),
        ("v. คำสั่งลบข้อมูลใน table",
         "DELETE FROM students WHERE name = 'jack';",
         "Deleted 1 record(s) successfully.\nRemaining Rows:\n(2, 'chompoo', 20, 'พยาบาล')\n(1, 'chokun', 23, 'วิศวะเอไอ')"),
        ("vi. คำสั่งลบ table",
         "DROP TABLE students;",
         "Table 'students' dropped successfully.")
    ]
    
    for title, query, output in sql_steps:
        doc.add_paragraph(title, style='ReportH2')
        t = doc.add_table(rows=2, cols=1)
        t.style = 'Table Grid'
        
        # Query
        cell_q = t.rows[0].cells[0]
        set_cell_background(cell_q, "E2E8F0")
        p_q = cell_q.paragraphs[0]
        p_q.text = "SQL Query:\n" + query.strip()
        p_q.style.font.name = 'Consolas'
        p_q.style.font.size = Pt(10)
        
        # Output
        cell_o = t.rows[1].cells[0]
        set_cell_background(cell_o, "F8FAFC")
        p_o = cell_o.paragraphs[0]
        p_o.text = "Output Log:\n" + output.strip()
        p_o.style.font.name = 'Consolas'
        p_o.style.font.size = Pt(10)
        doc.add_paragraph() # Spacing
    try:
        doc.add_picture(r"C:\Users\kt856\Pictures\Screenshots\Screenshot 2026-07-10 155353.png", width=Inches(5.8))
        p_caption1 = doc.add_paragraph("รูปที่ 1: ผลการทดสอบสคริปต์ SQL CRUD ของตาราง students")
        p_caption1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption1.style.font.italic = True
    except Exception as e:
        doc.add_paragraph(f"[ไม่สามารถโหลดรูปภาพ SQL CRUD ได้: {e}]").style.font.color.rgb = RGBColor(220, 38, 38)
    doc.add_paragraph()

    
    # 3. Section 2: Label Studio
    doc.add_paragraph('2. การใช้งาน Label Studio ร่วมกับ PostgreSQL', style='ReportH1')
    doc.add_paragraph('2.1 การตั้งค่าใน compose.yml เพื่อเชื่อมต่อ Label Studio กับ PostgreSQL')
    doc.add_paragraph('เราได้กำหนด Environment Variables ให้แก่บริการ label-studio เพื่อสั่งการให้ใช้ฐานข้อมูล PostgreSQL เป็น Backend (พอร์ตภายใน 5432) ดังนี้:')
    
    table_ls = doc.add_table(rows=1, cols=1)
    table_ls.style = 'Light Shading Accent 1'
    cell_ls = table_ls.rows[0].cells[0]
    set_cell_background(cell_ls, "F1F5F9")
    ls_config = """  label-studio:
    image: heartexlabs/label-studio:latest
    container_name: label_studio
    restart: always
    ports:
      - "8080:8080"
    environment:
      - DJANGO_DB=default
      - POSTGRE_NAME=student_db
      - POSTGRE_USER=postgres
      - POSTGRE_PASSWORD=mysecretpassword
      - POSTGRE_PORT=5432
      - POSTGRE_HOST=db
    depends_on:
      - db"""
    cell_ls.paragraphs[0].text = ls_config
    cell_ls.paragraphs[0].style.font.name = 'Consolas'
    cell_ls.paragraphs[0].style.font.size = Pt(10)
    
    doc.add_paragraph('\n2.2 ขั้นตอนการเปิดใช้งานและการสร้างบัญชีเข้าสู่ระบบ')
    doc.add_paragraph('เมื่อรันบริการด้วย Docker Compose เสร็จแล้ว สามารถเข้าสู่ระบบผ่านเว็บเบราว์เซอร์ที่ลิงก์ http://localhost:8080 และทำการลงทะเบียนบัญชีผู้ใช้ใหม่ (Sign up) จากนั้นเข้าสู่หน้าแดชบอร์ดหลักของระบบ')
    try:
        doc.add_picture(r"C:\Users\kt856\Pictures\Screenshots\Screenshot 2025-11-14 010557.png", width=Inches(5.8))
        p_caption2 = doc.add_paragraph("รูปที่ 2: หน้าต่างเข้าสู่ระบบ (Login / Sign Up) ของ Label Studio")
        p_caption2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption2.style.font.italic = True
        
        doc.add_paragraph() # Spacing
        
        doc.add_picture(r"C:\Users\kt856\Pictures\Screenshots\Screenshot 2026-07-10 154509.png", width=Inches(5.8))
        p_caption3 = doc.add_paragraph("รูปที่ 3: หน้าจอหลัก Dashboard หลังเข้าสู่ระบบของ Label Studio สำเร็จ")
        p_caption3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption3.style.font.italic = True
    except Exception as e:
        doc.add_paragraph(f"[ไม่สามารถโหลดรูปภาพหน้าจอ Label Studio ได้: {e}]").style.font.color.rgb = RGBColor(220, 38, 38)
    doc.add_paragraph()

    
    # 4. Section 3: Docker Containers
    doc.add_paragraph('3. สถานะของ Docker Containers ทั้งหมดในระบบ', style='ReportH1')
    doc.add_paragraph('คอนเทนเนอร์ทั้งหมดในโปรเจกต์นี้ประกอบด้วย redis (รันบนพอร์ต 6379), postgres_db (รันบนพอร์ต 5433 บนเครื่อง Host) และ label_studio (รันบนพอร์ต 8080) ทำงานร่วมกันได้อย่างเป็นระบบภายใต้เครือข่ายภายในเดียวกัน')
    try:
        doc.add_picture(r"C:\Users\kt856\Pictures\Screenshots\Screenshot 2026-07-10 155345.png", width=Inches(5.8))
        p_caption4 = doc.add_paragraph("รูปที่ 4: แสดงรายการ Docker Containers ทั้งหมดในระบบ (Redis, PostgreSQL, Label Studio)")
        p_caption4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption4.style.font.italic = True
    except Exception as e:
        doc.add_paragraph(f"[ไม่สามารถโหลดรูปภาพ Docker Container ได้: {e}]").style.font.color.rgb = RGBColor(220, 38, 38)
    doc.add_paragraph()

    
    # 5. Section 4: System Diagram
    doc.add_paragraph('4. ไดอะแกรมโครงสร้างระบบใหม่ (System Diagram)', style='ReportH1')
    doc.add_paragraph('ไดอะแกรมการทำงานและการแลกเปลี่ยนข้อมูลของระบบ:')
    
    # Draw simple table structure diagram
    diag_table = doc.add_table(rows=3, cols=3)
    diag_table.style = 'Table Grid'
    
    # Row 1
    cell_00 = diag_table.rows[0].cells[0]
    cell_00.paragraphs[0].text = "[ User Browser ]\n(localhost:8080)"
    set_cell_background(cell_00, "EFF6FF") # Light Blue
    
    diag_table.rows[0].cells[1].paragraphs[0].text = " ── HTTP ──> "
    
    cell_02 = diag_table.rows[0].cells[2]
    cell_02.paragraphs[0].text = "[ Label Studio Container ]\n(Port 8080)"
    set_cell_background(cell_02, "F0FDF4") # Light Green
    
    # Row 2
    diag_table.rows[1].cells[0].paragraphs[0].text = "[ Redis Container ]\n(Port 6379)"
    set_cell_background(diag_table.rows[1].cells[0], "FEF3C7") # Yellow
    
    diag_table.rows[1].cells[2].paragraphs[0].text = "       │\n       │ DB Connection\n       ▼"
    
    # Row 3
    cell_20 = diag_table.rows[2].cells[0]
    cell_20.paragraphs[0].text = "[ Named Volume ]\npostgresql-data & redis-data\n(Persistent storage)"
    set_cell_background(cell_20, "F5F5F5") # Grey
    
    diag_table.rows[2].cells[1].paragraphs[0].text = " <── Mount ── "
    
    cell_22 = diag_table.rows[2].cells[2]
    cell_22.paragraphs[0].text = "[ PostgreSQL Container ]\n(Port 5432 - Internal)\n(Port 5433 - Host)"
    set_cell_background(cell_22, "FEF2F2") # Light Red
    
    # Format all cells in diagram table to be centered
    for row in diag_table.rows:
        for c in row.cells:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraphs[0].style.font.name = 'Consolas'
            c.paragraphs[0].style.font.size = Pt(11)
            
    doc.add_paragraph('\n')
    
    # 6. Section 5: GitHub Link
    doc.add_paragraph('5. ลิงก์เก็บซอร์สโค้ด GitHub (GitHub Repository)', style='ReportH1')
    doc.add_paragraph('ซอร์สโค้ดสำหรับการรันระบบ Docker Compose, ชุดสคริปต์ทดสอบ และไดอะแกรมทั้งหมดถูกอัปโหลดขึ้น GitHub เรียบร้อยแล้วที่ลิงก์:')
    doc.add_paragraph('👉 ลิงก์โครงการ: https://github.com/chokunkie02/ai-ecosystem-workspace').style.font.bold = True
    
    # Save the document inside work1 directory
    current_dir = os.path.dirname(os.path.abspath(__file__))
    output_filename = os.path.join(current_dir, 'Assignment1_Submission_Checklist.docx')
    doc.save(output_filename)
    print(f"[Process] Report '{output_filename}' generated successfully.")

if __name__ == "__main__":
    generate_docx()
